"""最小验证：parser 返回的 page dict 是否带正确的 file_type。"""
import os
import tempfile

from services.parser import extract_pdf_pages, extract_txt_pages, extract_docx_pages


def test_txt():
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", delete=False, encoding="utf-8") as f:
        f.write("hello txt")
        path = f.name
    try:
        pages = extract_txt_pages(path)
        assert pages, "txt: 返回为空"
        assert pages[0]["file_type"] == "txt", f"txt: file_type={pages[0].get('file_type')}"
        print("✓ TXT  -> file_type='txt'")
    finally:
        os.unlink(path)


def test_docx():
    from docx import Document
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name
    doc = Document()
    doc.add_paragraph("hello docx")
    doc.save(path)
    try:
        pages = extract_docx_pages(path)
        assert pages, "docx: 返回为空"
        assert pages[0]["file_type"] == "docx", f"docx: file_type={pages[0].get('file_type')}"
        print("✓ DOCX -> file_type='docx'")
    finally:
        os.unlink(path)


def test_pdf():
    try:
        from fpdf import FPDF
    except ImportError:
        print("⚠ 跳过 PDF 测试（未安装 fpdf2，可 pip install fpdf2 后重试）")
        return
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="hello pdf")
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        path = f.name
    pdf.output(path)
    try:
        pages = extract_pdf_pages(path)
        assert pages, "pdf: 返回为空"
        assert pages[0]["file_type"] == "pdf", f"pdf: file_type={pages[0].get('file_type')}"
        print("✓ PDF  -> file_type='pdf'")
    finally:
        os.unlink(path)


if __name__ == "__main__":
    test_txt()
    test_docx()
    test_pdf()
    print("\n全部验证通过 ✓")
